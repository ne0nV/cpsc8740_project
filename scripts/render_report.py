from __future__ import annotations
# scripts/render_report.py
import os, json, datetime
from pathlib import Path

# Optional: generate a DOCX report from metrics.json.
# Requires: pip install python-docx

TEMPLATE = {
    "title": "AI Model & Performance Report",
    "subtitle": "Recommender System — Offline Evaluation",
    "author": os.getenv("REPORT_AUTHOR", "Todd Schuff"),
}

def main():
    try:
        from docx import Document
    except Exception as e:
        raise SystemExit("Missing dependency: python-docx. Install with `pip install python-docx`.") from e

    artifacts = Path("./artifacts")
    metrics_path = artifacts / "metrics.json"
    if not metrics_path.exists():
        raise SystemExit("artifacts/metrics.json not found. Run `python scripts/evaluate.py` first.")

    metrics = json.loads(metrics_path.read_text())

    doc = Document()

    # Title
    doc.add_heading(TEMPLATE["title"], 0)
    doc.add_paragraph(
        f"{TEMPLATE['subtitle']}\nAuthor: {TEMPLATE['author']}\nDate: {datetime.date.today().isoformat()}"
    )

    # 1. Overview
    doc.add_heading("1. Overview", level=1)
    doc.add_paragraph(
        f"""This report documents the baseline and model-based recommenders evaluated on an offline split.
Our primary objective is top-N recommendation quality using Precision@K, Recall@K, and NDCG@K with K = {metrics.get('k', 10)}.
We compare a popularity baseline against a LightFM matrix factorization model when available."""
    )

    # 2. Data
    doc.add_heading("2. Data", level=1)
    doc.add_paragraph(
        f"Dataset: {metrics['dataset']} | Rows: {metrics['n_rows']:,} | Users: {metrics['n_users']:,} | Items: {metrics['n_items']:,}."
    )
    doc.add_paragraph(
        "We use a timestamp-based leave-one-out split per user (last interaction to test)."
    )

    # 3. Methods
    doc.add_heading("3. Methods", level=1)
    for line in [
        "Popularity baseline: rank items by total observed ratings in training; filter out items already seen by the user.",
        "LightFM (if trained): hybrid matrix factorization with WARP loss on implicit feedback. Defaults: no_components=64, epochs=15, loss='warp'.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    # 4. Results
    doc.add_heading("4. Results (Precision@K / Recall@K / NDCG@K)", level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Model"
    hdr[1].text = "Precision@K"
    hdr[2].text = "Recall@K"
    hdr[3].text = "NDCG@K"

    for model_name, vals in metrics["models"].items():
        row = table.add_row().cells
        row[0].text = model_name
        row[1].text = str(vals.get("precision@k", ""))
        row[2].text = str(vals.get("recall@k", ""))
        row[3].text = str(vals.get("ndcg@k", ""))

    # 5. Error Analysis
    doc.add_heading("5. Error Analysis (Brief)", level=1)
    doc.add_paragraph(
        "Inspect users where Recall@K = 0 to understand failure modes (cold start, niche tastes). "
        "For popularity, errors often happen when a user’s last interaction was with a less-popular item. "
        "For LightFM, verify user/item ID maps and consider adding item side features."
    )

    # 6. Challenges & Next Steps
    doc.add_heading("6. Challenges & Next Steps", level=1)
    for line in [
        "Data sparsity and popularity bias can limit Recall@K. Consider blending in a content-based fallback for cold start.",
        "Tune LightFM hyperparameters (no_components, epochs, loss) via a validation split.",
        "For deployment, precompute top-N per user on a schedule and cache results to reduce latency.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    out_path = Path("./docs/AI_Model_Performance_Report.docx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Wrote {out_path}")

if __name__ == "__main__":
    main()
